"""
Test script for multi-document analysis endpoint.
Tests the complete workflow: upload ZIP -> normalize -> analyze
"""

import requests
import io
import zipfile
import openpyxl
from openpyxl import Workbook
import json

# Configuration
BASE_URL = "http://localhost:8001/api/v1/multi-document"

def create_sample_excel():
    """Create a sample T12 Excel file with expense data."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    
    # Add headers
    ws['A1'] = "ID"
    ws['B1'] = "Description"
    ws['C1'] = "Category"
    ws['D1'] = "Amount"
    ws['E1'] = "Date"
    ws['F1'] = "Notes"
    
    # Add sample expense data
    expenses = [
        (1, "Trash Removal Service", "Contract Services", 3600, "2024-01-15", "Monthly service"),
        (2, "Property Tax 1st Installment", "Real Estate Taxes", 12500, "2024-02-01", "Q1 payment"),
        (3, "PG&E Electric", "Utilities", 8400, "2024-01-20", "Annual total"),
        (4, "Water & Sewer", "Utilities", 4200, "2024-01-25", "Annual total"),
        (5, "Management Fee", "Management Fees", 15000, "2024-01-30", "Annual fee"),
        (6, "Plumbing Repairs", "Repairs & Maintenance", 2500, "2024-02-10", "Emergency repair"),
        (7, "HVAC Maintenance", "Repairs & Maintenance", 3200, "2024-03-01", "Annual service"),
        (8, "Landscaping Services", "Contract Services", 4800, "2024-01-15", "Monthly service"),
        (9, "Property Insurance", "Insurance", 9600, "2024-01-01", "Annual premium"),
    ]
    
    for idx, (id, desc, cat, amount, date, notes) in enumerate(expenses, start=2):
        ws[f'A{idx}'] = id
        ws[f'B{idx}'] = desc
        ws[f'C{idx}'] = cat
        ws[f'D{idx}'] = amount
        ws[f'E{idx}'] = date
        ws[f'F{idx}'] = notes
    
    # Save to bytes
    excel_bytes = io.BytesIO()
    wb.save(excel_bytes)
    excel_bytes.seek(0)
    return excel_bytes.read()


def create_test_zip():
    """Create a test ZIP file with the expected folder structure."""
    zip_buffer = io.BytesIO()
    
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
        # Create sample Excel file
        excel_content = create_sample_excel()
        
        # Add to Financials folder
        zip_file.writestr("04 - Financials/T12_Statement.xlsx", excel_content)
        
        # Add to Rent Roll folder
        zip_file.writestr("02 - Rent Roll/rent_roll.xlsx", excel_content)
    
    zip_buffer.seek(0)
    return zip_buffer.read()


def test_complete_workflow():
    """Test the complete workflow including analysis."""
    print("=" * 80)
    print("Testing Multi-Document Analysis Workflow")
    print("=" * 80)
    
    # Step 1: Create and upload ZIP package
    print("\n[STEP 1] Creating test ZIP package...")
    zip_content = create_test_zip()
    print(f"✓ Created ZIP with {len(zip_content)} bytes")
    
    print("\n[STEP 2] Uploading ZIP package...")
    files = {'file': ('Test_Property_Inputs.zip', zip_content, 'application/zip')}
    data = {'property_name': 'Test Property'}
    
    response = requests.post(f"{BASE_URL}/packages/upload-zip", files=files, data=data)
    
    if response.status_code != 200:
        print(f"✗ Upload failed: {response.status_code}")
        print(f"Error: {response.text}")
        return False
    
    package = response.json()
    package_id = package['package_id']
    print(f"✓ Package uploaded successfully")
    print(f"  Package ID: {package_id}")
    print(f"  Property: {package['property_name']}")
    print(f"  Documents: {sum(len(docs) for docs in package['documents'].values())} files")
    
    # Step 3: Normalize documents
    print("\n[STEP 3] Normalizing financial documents...")
    response = requests.post(f"{BASE_URL}/packages/{package_id}/normalize")
    
    if response.status_code != 200:
        print(f"✗ Normalization failed: {response.status_code}")
        print(f"Error: {response.text}")
        return False
    
    result = response.json()
    print(f"✓ Normalization completed")
    print(f"  Total items extracted: {result['total_items']}")
    print(f"  Average confidence: {result['confidence_average']:.2%}")
    
    # Step 4: Run analysis
    print("\n[STEP 4] Running financial analysis...")
    deal_params = {
        "growth_rate": 0.03,
        "exit_cap_rate": 0.06,
        "vacancy_rate": 0.03,
        "loan_amount": 5000000.0,
        "min_unit_count": 15,
        "max_unit_count": 80,
        "max_build_year": 1970
    }
    
    response = requests.post(
        f"{BASE_URL}/packages/{package_id}/analyze",
        json=deal_params
    )
    
    if response.status_code != 200:
        print(f"✗ Analysis failed: {response.status_code}")
        print(f"Error: {response.text}")
        return False
    
    analysis = response.json()
    print(f"✓ Analysis completed successfully")
    print(f"\n[ANALYSIS RESULTS]")
    print(f"  Status: {analysis['pass_fail_status']}")
    print(f"  Property: {analysis['property_meta']['address']}")
    print(f"  Units: {analysis['property_meta']['total_units']}")
    print(f"  Purchase Price: ${analysis['property_meta']['purchase_price']:,.2f}")
    print(f"\n[FINANCIAL METRICS]")
    print(f"  Historical NOI: ${analysis.get('historical_noi', 0):,.2f}")
    print(f"  Pro Forma NOI: ${analysis.get('pro_forma_noi', 0):,.2f}")
    print(f"  Pro Forma Expenses: ${analysis.get('pro_forma_expenses', 0):,.2f}")
    print(f"  Cap Rate: {analysis.get('cap_rate', 0):.2%}")
    print(f"  Exit Cap Rate: {analysis.get('exit_cap_rate', 0):.2%}")
    
    if analysis.get('gating_reasons'):
        print(f"\n[GATING REASONS]")
        for reason in analysis['gating_reasons']:
            print(f"  - {reason}")
    
    print("\n" + "=" * 80)
    print("✓ Complete workflow test PASSED!")
    print("=" * 80)
    
    return True


if __name__ == "__main__":
    try:
        success = test_complete_workflow()
        exit(0 if success else 1)
    except Exception as e:
        print(f"\n✗ Test failed with exception: {str(e)}")
        import traceback
        traceback.print_exc()
        exit(1)
