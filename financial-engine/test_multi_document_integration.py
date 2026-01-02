"""
Test script for multi-document extraction integration.
Tests the full workflow: upload ZIP -> normalize -> verify results
"""

import requests
import io
import zipfile
import openpyxl
from openpyxl import Workbook

# Configuration
BASE_URL = "http://localhost:8001/api/v1/multi-document"

def create_sample_excel():
    """Create a sample T12 Excel file with expense data."""
    wb = Workbook()
    ws = wb.active
    ws.title = "T12 Income Statement"
    
    # Add headers
    ws['A1'] = "Expense Category"
    ws['B1'] = "Annual Amount"
    
    # Add sample expense data
    expenses = [
        ("Trash Removal Service", 3600),
        ("Property Tax 1st Installment", 12500),
        ("PG&E Electric", 8400),
        ("Water & Sewer", 4200),
        ("Management Fee - ABC Property Mgmt", 15000),
        ("R&M - Plumbing", 2500),
        ("R&M - HVAC", 3200),
        ("Landscaping Services", 4800),
        ("Property Insurance", 9600),
    ]
    
    for idx, (name, amount) in enumerate(expenses, start=2):
        ws[f'A{idx}'] = name
        ws[f'B{idx}'] = amount
    
    # Save to bytes
    excel_bytes = io.BytesIO()
    wb.save(excel_bytes)
    excel_bytes.seek(0)
    return excel_bytes.read()


def create_test_zip():
    """Create a test ZIP file with the expected folder structure."""
    zip_buffer = io.BytesIO()
    
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
        # Create sample Excel file
        excel_content = create_sample_excel()
        
        # Add to Financials folder
        zip_file.writestr("04 - Financials/T12_Income_Statement.xlsx", excel_content)
        
        # Add a second file to Financials
        zip_file.writestr("04 - Financials/T12_P&L.xlsx", excel_content)
    
    zip_buffer.seek(0)
    return zip_buffer.read()


def test_upload_and_normalize():
    """Test the complete workflow."""
    print("=" * 60)
    print("Testing Multi-Document Extraction Integration")
    print("=" * 60)
    
    # Step 1: Create and upload ZIP package
    print("\n1. Creating test ZIP package...")
    zip_content = create_test_zip()
    print(f"   ✓ Created ZIP with {len(zip_content)} bytes")
    
    print("\n2. Uploading ZIP package...")
    files = {'file': ('Test_Property_Inputs.zip', zip_content, 'application/zip')}
    data = {'property_name': 'Test Property'}
    
    response = requests.post(f"{BASE_URL}/packages/upload-zip", files=files, data=data)
    
    if response.status_code != 200:
        print(f"   ✗ Upload failed: {response.status_code}")
        print(f"   Error: {response.text}")
        return False
    
    package = response.json()
    package_id = package['package_id']
    print(f"   ✓ Package uploaded successfully")
    print(f"   Package ID: {package_id}")
    print(f"   Property: {package['property_name']}")
    print(f"   Documents: {sum(len(docs) for docs in package['documents'].values())} files")
    
    # Step 3: Normalize documents
    print("\n3. Normalizing financial documents...")
    response = requests.post(f"{BASE_URL}/packages/{package_id}/normalize")
    
    if response.status_code != 200:
        print(f"   ✗ Normalization failed: {response.status_code}")
        print(f"   Error: {response.text}")
        return False
    
    result = response.json()
    print(f"   ✓ Normalization completed")
    print(f"   Total items extracted: {result['total_items']}")
    print(f"   Average confidence: {result['confidence_average']:.2%}")
    
    # Step 4: Display normalized items
    print("\n4. Normalized Expense Items:")
    print("-" * 60)
    
    for item in result['normalized_items']:
        confidence_bar = "█" * int(item['confidence'] * 10)
        print(f"\n   Raw Text: {item['raw_text']}")
        print(f"   Category: {item['normalized_value']}")
        print(f"   Confidence: {confidence_bar} {item['confidence']:.1%}")
        print(f"   Source: {item['source_document']}")
        if item.get('metadata', {}).get('amount'):
            print(f"   Amount: ${item['metadata']['amount']:,.2f}")
    
    print("\n" + "=" * 60)
    print("✓ Integration test completed successfully!")
    print("=" * 60)
    
    return True


if __name__ == "__main__":
    try:
        success = test_upload_and_normalize()
        exit(0 if success else 1)
    except Exception as e:
        print(f"\n✗ Test failed with exception: {str(e)}")
        import traceback
        traceback.print_exc()
        exit(1)
